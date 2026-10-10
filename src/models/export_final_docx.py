"""
Build publication-grade Word Document (.docx) from Team11_Phase2_IT5006_AY2627Sem1_final.md.

Ensures:
- 100% of text, data, tables, formulas, and high-resolution figures are embedded directly.
- No external links to charts or tables are used; all 15 figures and 17 tables are native Word elements.
- Professional executive styling, consistent typography, dynamic column widths, and clean pagination.
"""

import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/bensonwu/Projects/IT5006_GRP11')
SOURCE = ROOT / 'docs/reports/Team11_Phase2_IT5006_AY2627Sem1_final.md'
OUTPUT = ROOT / 'docs/reports/Team11_Phase2_IT5006_AY2627Sem1_final.docx'

MAX_CONTENT_WIDTH = 6.97  # A4 width (8.27) - 2 * 0.65 margin

def add_hyperlink(p, label, url, font_size=Pt(11), font_name='Times New Roman', is_bold=False):
    """Add a clickable hyperlink with professional styling to paragraph."""
    if not re.match(r'https?://', url):
        target_path = (SOURCE.parent / url).resolve()
        url = target_path.as_uri() if target_path.exists() else url
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')

    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')  # Professional hyperlink blue
    props.append(color)

    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)

    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:ascii'), font_name)
    rfonts.set(qn('w:hAnsi'), font_name)
    props.append(rfonts)

    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(font_size.pt * 2)))
    props.append(sz)

    if is_bold:
        b = OxmlElement('w:b')
        props.append(b)

    run.append(props)
    node = OxmlElement('w:t')
    node.text = label
    run.append(node)
    hyperlink.append(run)
    p._p.append(hyperlink)

def clean_latex(text):
    """Clean LaTeX math markup for publication-quality rendering in Word."""
    if not text:
        return ""
    res = text

    # Strip \text{...} and \mathbf{...} first to prevent nested braces breaking patterns
    res = re.sub(r'\\text\{([^}]*)\}', r'\1', res)
    res = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', res)
    res = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', res)
    res = res.replace(r'\_', '_')

    # Mathematical and Greek replacements
    replacements = [
        (r'\mathcal{I}', 'I'),
        (r'\mathcal{D}', 'D'),
        (r'\mathbb{I}', 'I'),
        (r'\mathbb{R}^+', 'R⁺'),
        (r'\tau^*', 'τ*'),
        (r'\tau', 'τ'),
        (r'\alpha', 'α'),
        (r'\beta', 'β'),
        (r'\kappa', 'κ'),
        (r'\sigma', 'σ'),
        (r'\mu', 'μ'),
        (r'\Delta', 'Δ'),
        (r'\le', '≤'),
        (r'\ge', '≥'),
        (r'\ne', '≠'),
        (r'\approx', '≈'),
        (r'\pm', '±'),
        (r'\times', '×'),
        (r'\cdot', '·'),
        (r'\to', '→'),
        (r'\in', '∈'),
        (r'\land', '∧'),
        (r'\lor', '∨'),
        (r'\mid', '|'),
        (r'R^2', 'R²'),
        (r'C_{FN}', 'C_FN'),
        (r'C_{FP}', 'C_FP'),
        (r'_{known}', '_known'),
        (r'_{checkout}', '_checkout'),
        (r'_{event}', '_event'),
        (r'\{', '{'),
        (r'\}', '}'),
        (r'$$', ''),
    ]
    for old, new in replacements:
        res = res.replace(old, new)

    # Clean subscripts like _{...} -> _...
    res = re.sub(r'_\{([^}]+)\}', r'_\1', res)
    # Strip $ delimiters
    res = re.sub(r'\$([^$\n]+)\$', r'\1', res)
    res = res.replace('$', '')
    return res

def inline(p, text, is_header=False, font_size=Pt(11), font_name='Times New Roman', is_italic=False, color_rgb=None):
    """Parse markdown inline formatting (bold, italic, code, links, math) and append runs to paragraph."""
    # Tokenize links, bold-italic, bold, italic, code, math
    pattern = r'(\[[^\]]+\]\([^\)]+\)|`[^`]+`|\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|\$[^$]+\$)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        link_m = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', token)
        if link_m:
            label, url = link_m.groups()
            add_hyperlink(p, clean_latex(label), url, font_size=font_size, font_name=font_name, is_bold=is_header)
            continue

        token_is_code = token.startswith('`') and token.endswith('`')
        token_is_bold_italic = token.startswith('***') and token.endswith('***')
        token_is_bold = token.startswith('**') and token.endswith('**')
        token_is_italic = token.startswith('*') and token.endswith('*') and not token_is_bold
        token_is_math = token.startswith('$') and token.endswith('$') and len(token) > 1

        if token_is_code:
            content = token[1:-1]
        elif token_is_bold_italic:
            content = clean_latex(token[3:-3])
        elif token_is_bold:
            content = clean_latex(token[2:-2])
        elif token_is_italic:
            content = clean_latex(token[1:-1])
        elif token_is_math:
            content = clean_latex(token[1:-1])
        else:
            content = clean_latex(token)

        run = p.add_run(content)
        run.font.name = font_name
        run.font.size = font_size

        if is_header or token_is_bold or token_is_bold_italic:
            run.bold = True
        if is_italic or token_is_italic or token_is_bold_italic:
            run.italic = True
        if token_is_math:
            run.font.name = 'Cambria Math'
        if token_is_code:
            run.font.name = 'Consolas'
            run.font.size = Pt(font_size.pt - 0.5)
            run.font.color.rgb = RGBColor(0x8A, 0x1E, 0x38)
        elif color_rgb:
            run.font.color.rgb = color_rgb

def calculate_column_widths(rows, max_width=MAX_CONTENT_WIDTH):
    """Calculate balanced column widths proportional to content lengths."""
    cols_count = len(rows[0])
    max_lens = [max(len(clean_latex(str(r[c]))) for r in rows) for c in range(cols_count)]
    weights = [max(l, 4) for l in max_lens]
    damped = [w ** 0.65 for w in weights]
    total_damped = sum(damped)
    widths = [(w / total_damped) * max_width for w in damped]
    min_w = 0.45
    for idx in range(len(widths)):
        if widths[idx] < min_w:
            widths[idx] = min_w
    curr_total = sum(widths)
    widths = [(w / curr_total) * max_width for w in widths]
    return widths

def add_formatted_table(doc, lines):
    """Convert markdown table lines to a styled docx Table."""
    raw_rows = [[c.strip() for c in l.strip('|').split('|')] for l in lines]
    if not raw_rows or len(raw_rows) < 2:
        return
    header = raw_rows[0]
    data_rows = raw_rows[2:] if raw_rows[1][0].startswith('-') or raw_rows[1][0].startswith(':') else raw_rows[1:]
    rows = [header] + data_rows
    cols_count = len(header)

    t = doc.add_table(rows=len(rows), cols=cols_count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    widths = calculate_column_widths(rows)

    if cols_count <= 4:
        tbl_font_size = Pt(9.5)
    elif cols_count <= 6:
        tbl_font_size = Pt(8.8)
    else:
        tbl_font_size = Pt(8.0)

    # Set table borders
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side, sz, color, val in [
        ('top', '8', '1F4E79', 'single'),
        ('bottom', '8', '1F4E79', 'single'),
        ('left', '4', 'D9D9D9', 'single'),
        ('right', '4', 'D9D9D9', 'single'),
        ('insideH', '4', 'E2E8F0', 'single'),
        ('insideV', '4', 'E2E8F0', 'single')
    ]:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), val)
        node.set(qn('w:sz'), sz)
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)
        borders.append(node)
    tblPr.append(borders)

    for row_idx, row_data in enumerate(rows):
        is_header = (row_idx == 0)
        table_row = t.rows[row_idx]
        trPr = table_row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if is_header:
            trPr.append(OxmlElement('w:tblHeader'))

        for col_idx, text in enumerate(row_data):
            cell = table_row.cells[col_idx]
            width = widths[col_idx]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            tcPr = cell._tc.get_or_add_tcPr()
            tcW = OxmlElement('w:tcW')
            tcW.set(qn('w:w'), str(int(width * 1440)))
            tcW.set(qn('w:type'), 'dxa')
            tcPr.append(tcW)

            margin = OxmlElement('w:tcMar')
            for side, w_dxa in [('top', '60'), ('bottom', '60'), ('left', '80'), ('right', '80')]:
                node = OxmlElement(f'w:{side}')
                node.set(qn('w:w'), w_dxa)
                node.set(qn('w:type'), 'dxa')
                margin.append(node)
            tcPr.append(margin)

            if is_header:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'E8EEF5')
                tcPr.append(shd)
            elif row_idx % 2 == 1:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F8FAFC')
                tcPr.append(shd)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05

            cleaned_cell_text = clean_latex(text)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif any(c in cleaned_cell_text for c in ['%', '±', 'days', '0.', '1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.']):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            inline(p, text, is_header=is_header, font_size=tbl_font_size, font_name='Times New Roman')

def build_document():
    print(f"Reading markdown source from {SOURCE}...")
    raw_text = SOURCE.read_text(encoding='utf-8')
    text = re.sub(r'<!--(?! pagebreak -->).*?-->', '', raw_text, flags=re.S)
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]

    # Page setup: A4
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.40)
    section.footer_distance = Inches(0.40)
    section.different_first_page_header_footer = True

    styles = doc.styles
    for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Caption']:
        if style_name in styles:
            s = styles[style_name]
            s.font.name = 'Times New Roman'
            s.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            rfonts = s.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
                rfonts.attrib.pop(qn('w:' + attr), None)
            for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                rfonts.set(qn('w:' + attr), 'Times New Roman')

    styles['Normal'].font.size = Pt(11)
    styles['Normal'].paragraph_format.line_spacing = 1.15
    styles['Normal'].paragraph_format.space_before = Pt(0)
    styles['Normal'].paragraph_format.space_after = Pt(4)

    styles['Title'].font.size = Pt(21)
    styles['Title'].font.bold = True
    styles['Title'].font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    styles['Title'].paragraph_format.space_before = Pt(20)
    styles['Title'].paragraph_format.space_after = Pt(8)
    styles['Title'].paragraph_format.line_spacing = 1.15
    styles['Title'].paragraph_format.keep_with_next = True

    styles['Subtitle'].font.size = Pt(12.5)
    styles['Subtitle'].font.italic = True
    styles['Subtitle'].font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    styles['Subtitle'].paragraph_format.space_before = Pt(0)
    styles['Subtitle'].paragraph_format.space_after = Pt(10)
    styles['Subtitle'].paragraph_format.keep_with_next = True

    styles['Heading 1'].font.size = Pt(14)
    styles['Heading 1'].font.bold = True
    styles['Heading 1'].font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    styles['Heading 1'].paragraph_format.space_before = Pt(16)
    styles['Heading 1'].paragraph_format.space_after = Pt(5)
    styles['Heading 1'].paragraph_format.line_spacing = 1.15
    styles['Heading 1'].paragraph_format.keep_with_next = True

    styles['Heading 2'].font.size = Pt(12.5)
    styles['Heading 2'].font.bold = True
    styles['Heading 2'].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    styles['Heading 2'].paragraph_format.space_before = Pt(12)
    styles['Heading 2'].paragraph_format.space_after = Pt(4)
    styles['Heading 2'].paragraph_format.line_spacing = 1.15
    styles['Heading 2'].paragraph_format.keep_with_next = True

    styles['Heading 3'].font.size = Pt(11.5)
    styles['Heading 3'].font.bold = True
    styles['Heading 3'].font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    styles['Heading 3'].paragraph_format.space_before = Pt(9)
    styles['Heading 3'].paragraph_format.space_after = Pt(3)
    styles['Heading 3'].paragraph_format.line_spacing = 1.15
    styles['Heading 3'].paragraph_format.keep_with_next = True

    styles['Caption'].font.size = Pt(9.5)
    styles['Caption'].font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    styles['Caption'].paragraph_format.space_before = Pt(3)
    styles['Caption'].paragraph_format.space_after = Pt(8)
    styles['Caption'].paragraph_format.line_spacing = 1.1

    header_para = section.header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_para.add_run('IT5006 Milestone 2 Technical Report | Team 11 (SmartCommerce Analytics)')
    header_run.font.name = 'Times New Roman'
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_field = OxmlElement('w:fldSimple')
    footer_field.set(qn('w:instr'), 'PAGE')
    footer_para._p.append(footer_field)

    first_footer_para = section.first_page_footer.paragraphs[0]
    first_footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_field = OxmlElement('w:fldSimple')
    first_field.set(qn('w:instr'), 'PAGE')
    first_footer_para._p.append(first_field)

    doc.core_properties.title = 'IT5006 Milestone 2 Final Technical & Strategic Report'
    doc.core_properties.author = 'Team 11 (SmartCommerce Analytics)'
    doc.core_properties.subject = 'Order-Placement Fulfilment Lead Time and Detractor Risk Modeling'
    doc.core_properties.comments = 'Master compiled Word document containing all text, tables, formulas, and high-resolution charts.'

    i = 0
    pending_break = False
    title_seen = False
    subtitle_seen = False
    meta_seen = False

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line == '<!-- pagebreak -->':
            pending_break = True
            i += 1
            continue

        if line == '---':
            i += 1
            continue

        # Code block fence (bash or mermaid)
        if line.startswith('```'):
            code_lines = []
            lang = line[3:].strip().lower()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence

            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.25)
            p_code.paragraph_format.right_indent = Inches(0.25)
            p_code.paragraph_format.space_before = Pt(4)
            p_code.paragraph_format.space_after = Pt(6)
            p_code.paragraph_format.line_spacing = 1.05

            run_code = p_code.add_run('\n'.join(code_lines))
            run_code.font.name = 'Consolas'
            run_code.font.size = Pt(8.5)
            run_code.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            continue

        # Display math ($$...$$)
        if line.startswith('$$') and line.endswith('$$') and len(line) > 4:
            math_content = clean_latex(line[2:-2].strip())
            p_m = doc.add_paragraph()
            p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_m.paragraph_format.space_before = Pt(6)
            p_m.paragraph_format.space_after = Pt(6)
            p_m.paragraph_format.keep_with_next = True
            run_m = p_m.add_run(math_content)
            run_m.font.name = 'Cambria Math'
            run_m.font.size = Pt(11)
            run_m.bold = True
            i += 1
            continue

        # Table block
        if line.startswith('|'):
            group = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i].strip())
                i += 1
            add_formatted_table(doc, group)
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(2)
            sp.paragraph_format.space_after = Pt(2)
            sp.paragraph_format.line_spacing = Pt(2)
            continue

        # Image block
        img_m = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if img_m:
            alt, img_rel_path = img_m.groups()
            img_path = (SOURCE.parent / img_rel_path).resolve()
            if not img_path.exists():
                alt_path = SOURCE.parent / 'figures' / os.path.basename(img_rel_path)
                if alt_path.exists():
                    img_path = alt_path
                else:
                    raise FileNotFoundError(f'Image not found: {img_path}')

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.paragraph_format.keep_with_next = True

            if pending_break:
                p_img.paragraph_format.page_break_before = True
                pending_break = False

            shape = p_img.add_run().add_picture(str(img_path), width=Inches(6.80))
            shape._inline.docPr.set('descr', alt or 'Report Figure')
            i += 1
            continue

        # Caption block (Figure X.Y / Table X.Y)
        if line.startswith('Figure ') or line.startswith('Table ') or line.startswith('*Figure ') or line.startswith('*Table '):
            p_cap = doc.add_paragraph(style='Caption')
            p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cap.paragraph_format.keep_together = True
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)

            clean_line = line.strip('*')
            fig_match = re.match(r'^((Figure|Table)\s+[A-Za-z0-9.]+[:\.]?\s*)(.*)', clean_line)
            if fig_match:
                prefix = fig_match.group(1)
                rest = fig_match.group(3)
                run_fig = p_cap.add_run(prefix)
                run_fig.bold = True
                run_fig.font.name = 'Times New Roman'
                run_fig.font.size = Pt(9.5)
                inline(p_cap, rest, font_size=Pt(9.5), font_name='Times New Roman', is_italic=True)
            else:
                inline(p_cap, clean_line, font_size=Pt(9.5), font_name='Times New Roman', is_italic=True)

            i += 1
            continue

        # Blockquote block
        if line.startswith('>'):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                raw_q = lines[i].strip()[1:].strip()
                if raw_q.startswith('[!NOTE]') or raw_q.startswith('[!IMPORTANT]') or raw_q.startswith('[!TIP]'):
                    raw_q = raw_q.split(']', 1)[1].strip()
                quote_lines.append(raw_q)
                i += 1

            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.35)
            p_q.paragraph_format.right_indent = Inches(0.20)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(6)
            p_q.paragraph_format.line_spacing = 1.12

            q_text = ' '.join(quote_lines)
            inline(p_q, q_text, font_size=Pt(10), font_name='Times New Roman', is_italic=True, color_rgb=RGBColor(0x33, 0x41, 0x55))
            continue

        # Headings
        heading_m = re.match(r'^(#{1,5})\s+(.*)', line)
        if heading_m:
            level = len(heading_m.group(1))
            h_text = heading_m.group(2)
            
            # Start major numbered sections and Appendices on fresh pages
            if level == 2 and any(h_text.startswith(sec) for sec in ['1.', '2.', '3.', '4.', '5.', '6.', '7.', 'Appendix']):
                pending_break = True

            if level == 1:
                p_h = doc.add_paragraph(style='Title')
                title_seen = True
                font_sz = Pt(20)
            elif level == 2:
                p_h = doc.add_paragraph(style='Heading 1')
                font_sz = Pt(14)
            elif level == 3:
                p_h = doc.add_paragraph(style='Heading 2')
                font_sz = Pt(12.5)
            elif level == 4:
                p_h = doc.add_paragraph(style='Heading 3')
                font_sz = Pt(11.5)
            else:
                p_h = doc.add_paragraph(style='Heading 4')
                font_sz = Pt(10.5)

            if pending_break:
                p_h.paragraph_format.page_break_before = True
                pending_break = False

            inline(p_h, h_text, is_header=True, font_size=font_sz)
            i += 1
            continue

        # Subtitle and Author Metadata
        if title_seen and not subtitle_seen:
            p = doc.add_paragraph(style='Subtitle')
            inline(p, line, font_size=Pt(12.5), is_italic=True)
            subtitle_seen = True
            i += 1
            continue
        elif title_seen and subtitle_seen and not meta_seen:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            inline(p, line, font_size=Pt(10.5))
            meta_seen = True
            i += 1
            continue

        # Bullet list items
        bullet_m = re.match(r'^([*\-•])\s+(.*)', line)
        if bullet_m:
            item_text = bullet_m.group(2)
            p_b = doc.add_paragraph()
            p_b.paragraph_format.left_indent = Inches(0.25)
            p_b.paragraph_format.first_line_indent = Inches(-0.15)
            p_b.paragraph_format.space_before = Pt(0)
            p_b.paragraph_format.space_after = Pt(2.5)
            p_b.paragraph_format.line_spacing = 1.15

            run_bullet = p_b.add_run('• ')
            run_bullet.font.name = 'Arial'
            run_bullet.font.size = Pt(10)
            run_bullet.bold = True
            run_bullet.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

            inline(p_b, item_text, font_size=Pt(10.5))
            i += 1
            continue

        # Numbered list items
        num_m = re.match(r'^(\d+)\.\s+(.*)', line)
        if num_m:
            num_str = num_m.group(1)
            item_text = num_m.group(2)
            p_n = doc.add_paragraph()
            p_n.paragraph_format.left_indent = Inches(0.28)
            p_n.paragraph_format.first_line_indent = Inches(-0.18)
            p_n.paragraph_format.space_before = Pt(0)
            p_n.paragraph_format.space_after = Pt(2.5)
            p_n.paragraph_format.line_spacing = 1.15

            run_num = p_n.add_run(f'{num_str}. ')
            run_num.font.name = 'Times New Roman'
            run_num.font.size = Pt(10.5)
            run_num.bold = True
            run_num.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

            inline(p_n, item_text, font_size=Pt(10.5))
            i += 1
            continue

        # Normal body paragraph
        p = doc.add_paragraph(style='Normal')
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False

        if re.match(r'^\[\d+\]', line):
            p.paragraph_format.left_indent = Inches(0.30)
            p.paragraph_format.first_line_indent = Inches(-0.30)
            p.paragraph_format.space_after = Pt(3)

        inline(p, line, font_size=Pt(11))
        i += 1

    print(f"Saving compiled Word document to {OUTPUT}...")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"SUCCESS: Generated {OUTPUT}")
    print(f"File size: {OUTPUT.stat().st_size:,} bytes")

if __name__ == '__main__':
    build_document()
